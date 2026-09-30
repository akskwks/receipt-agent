from __future__ import annotations

import tempfile
import unittest
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from receipt_agent import (
    ReceiptAgent,
    ReceiptAgentError,
    extract_largest_amount,
    find_latest_month_folder,
    find_workbook,
    parse_receipt_filename,
    set_log_file,
    update_workbook,
)


class ReceiptAgentTests(unittest.TestCase):
    def temporary_directory(self):
        return tempfile.TemporaryDirectory(dir=Path.cwd())

    def test_parse_receipt_filename(self) -> None:
        info = parse_receipt_filename(Path("[주인공][9월]활동금액_김민준.jpg"))
        self.assertEqual(9, info.month)
        self.assertEqual("김민준", info.name)

    def test_rejects_invalid_name_length(self) -> None:
        with self.assertRaises(ReceiptAgentError):
            parse_receipt_filename(Path("[주인공][9월]활동금액_김민수현.jpg"))

    def test_extracts_largest_won_amount(self) -> None:
        texts = ["승인금액 20,000원", "-34,396원", "계좌 94320200496799"]
        self.assertEqual(34_396, extract_largest_amount(texts))

    def test_extracts_amount_when_comma_is_missing(self) -> None:
        self.assertEqual(34_396, extract_largest_amount(["-34396원"]))

    def test_ignores_unformatted_account_number(self) -> None:
        with self.assertRaises(ReceiptAgentError):
            extract_largest_amount(["국민 94320200496799"])

    def test_selects_highest_month_receipt_folder(self) -> None:
        with self.temporary_directory() as temp_dir:
            root = Path(temp_dir)
            (root / "09월_영수증").mkdir()
            (root / "10월_영수증").mkdir()
            (root / "12월").mkdir()

            month, folder = find_latest_month_folder(root)
            self.assertEqual(10, month)
            self.assertEqual("10월_영수증", folder.name)

    def test_ignores_invalid_month_receipt_folder(self) -> None:
        with self.temporary_directory() as temp_dir:
            root = Path(temp_dir)
            (root / "00월_영수증").mkdir()
            (root / "13월_영수증").mkdir()
            (root / "09월").mkdir()

            with self.assertRaises(ReceiptAgentError):
                find_latest_month_folder(root)

    def test_agent_switches_when_higher_month_folder_is_added(self) -> None:
        with self.temporary_directory() as temp_dir:
            root = Path(temp_dir)
            september = root / "09월_영수증"
            october = root / "10월_영수증"
            september.mkdir()
            agent = ReceiptAgent(
                root=root,
                year=2026,
                interval=0.01,
                dry_run=True,
                excel_backend="openpyxl",
            )

            try:
                self.assertTrue(agent.refresh_target_folder())
                self.assertEqual(september, agent.folder)
                self.assertEqual(
                    september / "logs" / ".receipt_agent_state.json",
                    agent.state_path,
                )
                self.assertTrue(agent.state_path.exists())
                self.assertTrue((september / "logs" / "receipt_agent.log").exists())

                october.mkdir()
                self.assertTrue(agent.refresh_target_folder())
                self.assertEqual(october, agent.folder)
                self.assertEqual(10, agent.target_month)
                self.assertEqual(
                    october / "logs" / ".receipt_agent_state.json",
                    agent.state_path,
                )
                self.assertTrue(agent.state_path.exists())
                self.assertTrue((october / "logs" / "receipt_agent.log").exists())
            finally:
                set_log_file(None)

    def test_legacy_state_file_moves_into_logs_folder(self) -> None:
        with self.temporary_directory() as temp_dir:
            root = Path(temp_dir)
            folder = root / "09월_영수증"
            folder.mkdir()
            legacy_state = folder / ".receipt_agent_state.json"
            legacy_state.write_text(
                json.dumps({"receipt.jpg": "fingerprint"}), encoding="utf-8"
            )
            agent = ReceiptAgent(
                root=root,
                year=2026,
                interval=0.01,
                dry_run=True,
                excel_backend="openpyxl",
            )

            try:
                self.assertTrue(agent.refresh_target_folder())
                self.assertFalse(legacy_state.exists())
                self.assertTrue(
                    (folder / "logs" / ".receipt_agent_state.json").exists()
                )
                self.assertEqual("fingerprint", agent.state["receipt.jpg"])
            finally:
                set_log_file(None)

    def test_find_workbook_and_update_named_row(self) -> None:
        with self.temporary_directory() as temp_dir:
            folder = Path(temp_dir)
            workbook_path = (
                folder / "9월_2026_주인공_회원별_활동_금액_목록.xlsx"
            )
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "9월 멤버별 활동 금액"
            sheet.append(["순번", "이름", "활동유무", "지원금", "활동금액"])
            sheet.append([1, "김민준", "N", 20_000, 0])
            workbook.save(workbook_path)

            found = find_workbook(folder, 9, 2026)
            self.assertEqual(workbook_path, found)
            result = update_workbook(
                found,
                "김민준",
                34_396,
                backend="openpyxl",
            )
            self.assertEqual("E2", result.cell_coordinate)
            self.assertEqual(0, result.previous_value)

            updated = load_workbook(workbook_path, data_only=False)
            self.assertEqual(34_396, updated.active["E2"].value)
            updated.close()
            self.assertEqual([], list(folder.glob("*.backup-*.xlsx")))

    def test_sample_suffix_workbook_is_not_selected(self) -> None:
        with self.temporary_directory() as temp_dir:
            folder = Path(temp_dir)
            workbook = Workbook()
            workbook.save(
                folder / "9월_2026_주인공_회원별_활동_금액_목록_샘플.xlsx"
            )

            with self.assertRaises(ReceiptAgentError):
                find_workbook(folder, 9, 2026)

    def test_duplicate_name_is_rejected(self) -> None:
        with self.temporary_directory() as temp_dir:
            workbook_path = Path(temp_dir) / "test.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["이름", "활동금액"])
            sheet.append(["김민준", 0])
            sheet.append(["김민준", 0])
            workbook.save(workbook_path)

            with self.assertRaises(ReceiptAgentError):
                update_workbook(
                    workbook_path,
                    "김민준",
                    34_396,
                    backend="openpyxl",
                )


if __name__ == "__main__":
    unittest.main()
