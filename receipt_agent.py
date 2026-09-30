from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import re
import sys
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageEnhance, ImageOps


DEFAULT_ROOT = Path(
    r"C:\Users\THE KE\Desktop\mjkim\주인공(AI동호회)\2026"
)
DEFAULT_YEAR = 2026
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}
MONTH_FOLDER_RE = re.compile(r"^(?P<month>0?[1-9]|1[0-2])월_영수증$")
RECEIPT_FILENAME_RE = re.compile(
    r"^\[주인공\]\[(?P<month>\d{1,2})월\]활동금액_(?P<name>[가-힣]{3})$"
)
AMOUNT_RE = re.compile(
    r"(?<!\d)[\-\u2212\u2013\u2014]?\s*(?P<number>"
    r"(?:\d{1,3}(?:[,\s]\d{3})+)|(?:\d{4,})"
    r")\s*(?P<won>원)?(?!\d)"
)
LOG_FORMAT = "%(asctime)s %(levelname)s %(message)s"
LOG_DATE_FORMAT = "%Y-%m-%d %H:%M:%S"
_file_log_handler: logging.FileHandler | None = None


class ReceiptAgentError(RuntimeError):
    pass


@dataclass(frozen=True)
class ReceiptInfo:
    path: Path
    month: int
    name: str


@dataclass(frozen=True)
class UpdateResult:
    workbook_path: Path
    sheet_name: str
    cell_coordinate: str
    previous_value: object
    amount: int


def configure_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        handlers=[logging.StreamHandler(sys.stdout)],
        format=LOG_FORMAT,
        datefmt=LOG_DATE_FORMAT,
        force=True,
    )


def set_log_file(log_path: Path | None) -> None:
    global _file_log_handler
    root_logger = logging.getLogger()
    if _file_log_handler is not None:
        root_logger.removeHandler(_file_log_handler)
        _file_log_handler.close()
        _file_log_handler = None

    if log_path is None:
        return

    try:
        log_path.parent.mkdir(parents=True, exist_ok=True)
        handler = logging.FileHandler(log_path, encoding="utf-8")
        handler.setFormatter(logging.Formatter(LOG_FORMAT, datefmt=LOG_DATE_FORMAT))
        root_logger.addHandler(handler)
        _file_log_handler = handler
    except OSError as exc:
        logging.error("로그 파일을 열 수 없어 콘솔에만 기록합니다: %s", exc)


def parse_receipt_filename(path: Path) -> ReceiptInfo:
    if path.suffix.lower() not in IMAGE_EXTENSIONS:
        raise ReceiptAgentError(f"지원하지 않는 이미지 형식입니다: {path.name}")

    match = RECEIPT_FILENAME_RE.fullmatch(path.stem)
    if not match:
        raise ReceiptAgentError(
            "영수증 파일명이 규칙과 다릅니다: "
            f"{path.name} (예: [주인공][9월]활동금액_김민준.jpg)"
        )

    month = int(match.group("month"))
    if not 1 <= month <= 12:
        raise ReceiptAgentError(f"파일명의 월이 올바르지 않습니다: {month}월")
    return ReceiptInfo(path=path, month=month, name=match.group("name"))


def find_latest_month_folder(root: Path) -> tuple[int, Path]:
    month_folders: list[tuple[int, Path]] = []
    for path in root.iterdir():
        if not path.is_dir():
            continue
        match = MONTH_FOLDER_RE.fullmatch(path.name)
        if match:
            month_folders.append((int(match.group("month")), path))

    if not month_folders:
        raise ReceiptAgentError(
            f"월별 영수증 폴더를 찾지 못했습니다: {root} "
            "(예: 09월_영수증, 10월_영수증)"
        )

    latest_month = max(month for month, _path in month_folders)
    latest_paths = [path for month, path in month_folders if month == latest_month]
    if len(latest_paths) > 1:
        names = ", ".join(sorted(path.name for path in latest_paths))
        raise ReceiptAgentError(
            f"같은 월의 영수증 폴더가 여러 개입니다: {names}"
        )
    return latest_month, latest_paths[0]


def normalize_ocr_texts(texts: Iterable[str]) -> list[str]:
    normalized: list[str] = []
    for text in texts:
        value = str(text).strip()
        if value:
            normalized.append(value)
            normalized.append(value.replace("O", "0").replace("o", "0"))
    if normalized:
        normalized.append(" ".join(normalized[::2]))
    return normalized


def extract_largest_amount(texts: Iterable[str]) -> int:
    candidates: set[int] = set()
    for text in normalize_ocr_texts(texts):
        for match in AMOUNT_RE.finditer(text):
            raw_number = match.group("number")
            has_group_separator = bool(re.search(r"[,\s]", raw_number))
            has_won = bool(match.group("won"))
            if not has_group_separator and not has_won:
                continue

            amount = int(re.sub(r"\D", "", raw_number))
            if amount >= 10_000:
                candidates.add(amount)

    if not candidates:
        raise ReceiptAgentError(
            "이미지에서 10,000원 이상의 금액을 찾지 못했습니다. "
            "금액과 '원' 또는 천 단위 쉼표가 선명한지 확인해 주세요."
        )
    return max(candidates)


def _preprocess_image(image_path: Path, output_path: Path) -> None:
    with Image.open(image_path) as image:
        image = ImageOps.exif_transpose(image).convert("L")
        image = ImageOps.autocontrast(image)
        image = ImageEnhance.Contrast(image).enhance(1.4)
        if image.width < 1400:
            scale = 1400 / image.width
            image = image.resize(
                (round(image.width * scale), round(image.height * scale)),
                Image.Resampling.LANCZOS,
            )
        image.save(output_path, format="PNG")


class RapidOcrReader:
    def __init__(self) -> None:
        try:
            from rapidocr import RapidOCR
        except ImportError as exc:
            raise ReceiptAgentError(
                "OCR 패키지가 설치되지 않았습니다. "
                "'python -m pip install -r requirements.txt'를 실행해 주세요."
            ) from exc
        self._engine = RapidOCR()

    def read_texts(self, image_path: Path) -> list[str]:
        with tempfile.TemporaryDirectory(prefix="receipt-ocr-") as temp_dir:
            processed_path = Path(temp_dir) / "preprocessed.png"
            _preprocess_image(image_path, processed_path)
            result = self._engine(str(processed_path))
        texts = self._result_texts(result)
        if not texts:
            raise ReceiptAgentError(f"OCR 결과가 비어 있습니다: {image_path.name}")
        return texts

    @staticmethod
    def _result_texts(result: object) -> list[str]:
        for source in (result, getattr(result, "result", None)):
            if source is None:
                continue
            txts = getattr(source, "txts", None)
            if txts is not None:
                return [str(text) for text in txts]

        if isinstance(result, tuple) and result:
            result = result[0]
        texts: list[str] = []
        if isinstance(result, Sequence) and not isinstance(result, (str, bytes)):
            for item in result:
                if (
                    isinstance(item, Sequence)
                    and not isinstance(item, (str, bytes))
                    and len(item) >= 2
                ):
                    texts.append(str(item[1]))
        return texts


def find_workbook(folder: Path, month: int, year: int) -> Path:
    exact_name = f"{month}월_{year}_주인공_회원별_활동_금액_목록.xlsx"
    exact_path = folder / exact_name
    if exact_path.exists():
        return exact_path
    raise ReceiptAgentError(
        f"대상 엑셀 파일을 찾지 못했습니다: {exact_name} ({folder})"
    )


def _normalize_cell_text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _find_target_cell(workbook: object, name: str):
    matches: list[tuple[object, object]] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            headers = {_normalize_cell_text(cell.value): cell.column for cell in row}
            if "이름" not in headers or "활동금액" not in headers:
                continue

            name_column = headers["이름"]
            amount_column = headers["활동금액"]
            for row_number in range(row[0].row + 1, worksheet.max_row + 1):
                if _normalize_cell_text(
                    worksheet.cell(row=row_number, column=name_column).value
                ) == name:
                    matches.append(
                        (
                            worksheet,
                            worksheet.cell(row=row_number, column=amount_column),
                        )
                    )
            break

    if not matches:
        raise ReceiptAgentError(f"엑셀에서 이름 '{name}'을 찾지 못했습니다.")
    if len(matches) > 1:
        locations = ", ".join(
            f"{worksheet.title}!{cell.coordinate}" for worksheet, cell in matches
        )
        raise ReceiptAgentError(f"엑셀에서 이름 '{name}'이 중복됩니다: {locations}")
    return matches[0]


def _lock_file_for(workbook_path: Path) -> Path:
    return workbook_path.with_name(f"~${workbook_path.name}")


def update_workbook(
    workbook_path: Path,
    name: str,
    amount: int,
    *,
    dry_run: bool = False,
    backend: str = "excel",
) -> UpdateResult:
    if _lock_file_for(workbook_path).exists() and not dry_run:
        raise ReceiptAgentError(
            f"Excel에서 파일이 열려 있습니다. 닫은 뒤 다시 시도해 주세요: {workbook_path.name}"
        )

    if dry_run or backend == "openpyxl":
        return _update_workbook_with_openpyxl(
            workbook_path,
            name,
            amount,
            dry_run=dry_run,
        )
    if backend != "excel":
        raise ReceiptAgentError(f"지원하지 않는 Excel 저장 방식입니다: {backend}")
    return _update_workbook_with_excel(workbook_path, name, amount)


def _update_workbook_with_openpyxl(
    workbook_path: Path,
    name: str,
    amount: int,
    *,
    dry_run: bool,
) -> UpdateResult:
    workbook = load_workbook(workbook_path, data_only=False, keep_links=True)
    worksheet, target_cell = _find_target_cell(workbook, name)
    previous_value = target_cell.value
    target_cell.value = amount

    if not dry_run:
        temp_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=f".{workbook_path.stem}-",
                suffix=workbook_path.suffix,
                dir=workbook_path.parent,
                delete=False,
            ) as temp_file:
                temp_path = Path(temp_file.name)
            workbook.save(temp_path)
            os.replace(temp_path, workbook_path)
        finally:
            workbook.close()
            if temp_path and temp_path.exists():
                temp_path.unlink()
    else:
        workbook.close()

    return UpdateResult(
        workbook_path=workbook_path,
        sheet_name=worksheet.title,
        cell_coordinate=target_cell.coordinate,
        previous_value=previous_value,
        amount=amount,
    )


def _update_workbook_with_excel(
    workbook_path: Path,
    name: str,
    amount: int,
) -> UpdateResult:
    if os.name != "nt":
        raise ReceiptAgentError(
            "Excel 저장 방식은 Windows에서만 사용할 수 있습니다. "
            "다른 운영체제에서는 --excel-backend openpyxl을 사용해 주세요."
        )
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise ReceiptAgentError(
            "Excel 연동 패키지가 설치되지 않았습니다. "
            "'python -m pip install -r requirements.txt'를 실행해 주세요."
        ) from exc

    excel = None
    workbook = None
    pythoncom.CoInitialize()
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AskToUpdateLinks = False
        workbook = excel.Workbooks.Open(
            str(workbook_path),
            UpdateLinks=0,
            ReadOnly=False,
            IgnoreReadOnlyRecommended=True,
        )
        if workbook.ReadOnly:
            raise ReceiptAgentError(f"엑셀 파일이 읽기 전용입니다: {workbook_path.name}")

        matches: list[tuple[object, object, str]] = []
        for worksheet in workbook.Worksheets:
            used_range = worksheet.UsedRange
            first_row = used_range.Row
            first_column = used_range.Column
            last_row = first_row + used_range.Rows.Count - 1
            last_column = first_column + used_range.Columns.Count - 1

            for row_number in range(first_row, last_row + 1):
                headers: dict[str, int] = {}
                for column_number in range(first_column, last_column + 1):
                    value = worksheet.Cells(row_number, column_number).Value
                    headers[_normalize_cell_text(value)] = column_number
                if "이름" not in headers or "활동금액" not in headers:
                    continue

                name_column = headers["이름"]
                amount_column = headers["활동금액"]
                for data_row in range(row_number + 1, last_row + 1):
                    value = worksheet.Cells(data_row, name_column).Value
                    if _normalize_cell_text(value) == name:
                        matches.append(
                            (
                                worksheet,
                                worksheet.Cells(data_row, amount_column),
                                f"{get_column_letter(amount_column)}{data_row}",
                            )
                        )
                break

        if not matches:
            raise ReceiptAgentError(f"엑셀에서 이름 '{name}'을 찾지 못했습니다.")
        if len(matches) > 1:
            locations = ", ".join(
                f"{worksheet.Name}!{cell_coordinate}"
                for worksheet, _cell, cell_coordinate in matches
            )
            raise ReceiptAgentError(f"엑셀에서 이름 '{name}'이 중복됩니다: {locations}")

        worksheet, target_cell, cell_coordinate = matches[0]
        previous_value = target_cell.Value
        sheet_name = worksheet.Name
        target_cell.Value = amount
        workbook.Save()
        workbook.Close(SaveChanges=False)
        workbook = None
        return UpdateResult(
            workbook_path=workbook_path,
            sheet_name=sheet_name,
            cell_coordinate=cell_coordinate,
            previous_value=previous_value,
            amount=amount,
        )
    except ReceiptAgentError:
        raise
    except Exception as exc:
        raise ReceiptAgentError(f"Excel 저장 중 오류가 발생했습니다: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        if excel is not None:
            excel.Quit()
        pythoncom.CoUninitialize()


def file_fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_state(state_path: Path) -> dict[str, str]:
    if not state_path.exists():
        return {}
    try:
        data = json.loads(state_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        logging.warning("처리 이력 파일을 읽지 못해 새로 시작합니다: %s", exc)
        return {}
    return data if isinstance(data, dict) else {}


def save_state(state_path: Path, state: dict[str, str]) -> None:
    state_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = state_path.with_suffix(f"{state_path.suffix}.tmp")
    temp_path.write_text(
        json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    os.replace(temp_path, state_path)


def wait_until_stable(path: Path, interval: float = 0.5, checks: int = 3) -> bool:
    previous: tuple[int, int] | None = None
    stable_count = 0
    for _ in range(max(checks * 4, 4)):
        try:
            stat = path.stat()
            current = (stat.st_size, stat.st_mtime_ns)
        except FileNotFoundError:
            return False
        if current == previous and current[0] > 0:
            stable_count += 1
            if stable_count >= checks:
                return True
        else:
            stable_count = 0
        previous = current
        time.sleep(interval)
    return False


class ReceiptAgent:
    def __init__(
        self,
        root: Path,
        year: int,
        interval: float,
        dry_run: bool,
        excel_backend: str,
    ) -> None:
        self.root = root
        self.folder: Path | None = None
        self.target_month: int | None = None
        self.year = year
        self.interval = interval
        self.dry_run = dry_run
        self.excel_backend = excel_backend
        self.state_path: Path | None = None
        self.state: dict[str, str] = {}
        self._ocr: RapidOcrReader | None = None
        self._target_error: str | None = None

    @property
    def ocr(self) -> RapidOcrReader:
        if self._ocr is None:
            logging.info("OCR 엔진을 초기화합니다.")
            self._ocr = RapidOcrReader()
        return self._ocr

    def refresh_target_folder(self) -> bool:
        try:
            month, folder = find_latest_month_folder(self.root)
        except (OSError, ReceiptAgentError) as exc:
            message = str(exc)
            if message != self._target_error:
                logging.error("처리 대상 폴더 선택 실패: %s", message)
                self._target_error = message
            self.folder = None
            self.target_month = None
            self.state_path = None
            self.state = {}
            return False

        self._target_error = None
        if folder != self.folder:
            try:
                logs_dir = folder / "logs"
                logs_dir.mkdir(parents=True, exist_ok=True)

                legacy_log_path = folder / "receipt_agent.log"
                log_path = logs_dir / "receipt_agent.log"
                if legacy_log_path.exists() and not log_path.exists():
                    os.replace(legacy_log_path, log_path)

                legacy_state_path = folder / ".receipt_agent_state.json"
                state_path = logs_dir / ".receipt_agent_state.json"
                if legacy_state_path.exists() and not state_path.exists():
                    os.replace(legacy_state_path, state_path)
                if not state_path.exists():
                    save_state(state_path, {})
            except OSError as exc:
                logging.error("월별 로그 폴더 준비 실패: %s", exc)
                return False

            set_log_file(log_path)
            self.folder = folder
            self.target_month = month
            self.state_path = state_path
            self.state = load_state(state_path)
            logging.info("처리 대상 폴더를 %s로 설정했습니다.", folder)
        return True

    def image_paths(self) -> list[Path]:
        if self.folder is None:
            return []
        return sorted(
            path
            for path in self.folder.iterdir()
            if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS
        )

    def process(self, image_path: Path) -> bool:
        try:
            receipt = parse_receipt_filename(image_path)
            if self.target_month is None or receipt.month != self.target_month:
                raise ReceiptAgentError(
                    f"파일명의 월({receipt.month}월)과 대상 폴더의 월"
                    f"({self.target_month}월)이 다릅니다."
                )
            if not wait_until_stable(image_path):
                raise ReceiptAgentError(f"파일 저장이 완료되지 않았습니다: {image_path.name}")
            fingerprint = file_fingerprint(image_path)
            state_key = image_path.name
            already_processed = self.state.get(state_key) == fingerprint or any(
                Path(saved_key).name == image_path.name
                and saved_fingerprint == fingerprint
                for saved_key, saved_fingerprint in self.state.items()
            )
            if already_processed:
                if not self.dry_run and self.state.get(state_key) != fingerprint:
                    self.state[state_key] = fingerprint
                    if self.state_path is not None:
                        save_state(self.state_path, self.state)
                return False

            texts = self.ocr.read_texts(image_path)
            amount = extract_largest_amount(texts)
            if self.folder is None:
                raise ReceiptAgentError("처리 대상 폴더가 선택되지 않았습니다.")
            workbook_path = find_workbook(self.folder, receipt.month, self.year)
            result = update_workbook(
                workbook_path,
                receipt.name,
                amount,
                dry_run=self.dry_run,
                backend=self.excel_backend,
            )
            logging.info(
                "%s -> %s!%s: %r -> %s원%s",
                image_path.name,
                result.sheet_name,
                result.cell_coordinate,
                result.previous_value,
                f"{result.amount:,}",
                " (미리보기)" if self.dry_run else "",
            )
            logging.info("OCR 인식 문자열: %s", " | ".join(texts))

            if not self.dry_run:
                self.state[state_key] = fingerprint
                if self.state_path is None:
                    raise ReceiptAgentError("처리 이력 경로가 설정되지 않았습니다.")
                save_state(self.state_path, self.state)
            return True
        except ReceiptAgentError as exc:
            logging.error("%s 처리 실패: %s", image_path.name, exc)
        except Exception:
            logging.exception("%s 처리 중 예상하지 못한 오류가 발생했습니다.", image_path.name)
        return False

    def run_once(self) -> int:
        if not self.refresh_target_folder():
            return 0
        processed = 0
        for image_path in self.image_paths():
            processed += int(self.process(image_path))
        return processed

    def watch(self) -> None:
        logging.info("월별 폴더 감시를 시작합니다: %s", self.root)
        logging.info("종료하려면 Ctrl+C를 누르세요.")
        try:
            while True:
                self.run_once()
                time.sleep(self.interval)
        except KeyboardInterrupt:
            logging.info("폴더 감시를 종료합니다.")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="영수증 이미지 금액을 OCR로 읽어 회원별 활동금액 엑셀에 반영합니다."
    )
    parser.add_argument(
        "--root",
        "--folder",
        dest="root",
        type=Path,
        default=DEFAULT_ROOT,
        help="월별 NN월_영수증 폴더가 들어 있는 연도 루트 경로",
    )
    parser.add_argument("--year", type=int, default=DEFAULT_YEAR, help="엑셀 파일의 연도")
    parser.add_argument(
        "--interval", type=float, default=2.0, help="폴더 확인 간격(초)"
    )
    parser.add_argument(
        "--once", action="store_true", help="현재 이미지들을 한 번 처리하고 종료"
    )
    parser.add_argument(
        "--dry-run", action="store_true", help="엑셀을 저장하지 않고 결과만 확인"
    )
    parser.add_argument(
        "--excel-backend",
        choices=("excel", "openpyxl"),
        default="excel",
        help="엑셀 저장 방식(기본: Excel 프로그램을 이용해 원본 구조 보존)",
    )
    return parser


def main() -> int:
    args = build_parser().parse_args()
    root = args.root.expanduser().resolve()
    configure_logging()
    if not root.is_dir():
        logging.error("연도 루트 폴더가 존재하지 않습니다: %s", root)
        return 2
    if args.interval <= 0:
        logging.error("확인 간격은 0보다 커야 합니다.")
        return 2

    agent = ReceiptAgent(
        root=root,
        year=args.year,
        interval=args.interval,
        dry_run=args.dry_run,
        excel_backend=args.excel_backend,
    )
    if args.once:
        count = agent.run_once()
        logging.info("처리 완료: %d개", count)
    else:
        agent.watch()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
